import pandas as pd
import os
import yt_dlp as yt
import openpyxl
from openpyxl import load_workbook, Workbook
from openpyxl.utils import get_column_letter
from mutagen.easyid3 import EasyID3

excel = "./music.xlsx"
archivo = load_workbook(excel)
ruta_base = "C:/Users/maico/Music/music"
albums = []

url = "https://www.youtube.com/watch?v=DCQHfrqMKXA"

def obtener_nombre_album(url):
    ydl_opts = {
        'quiet': True,
        'skip_download': True,
        'extract_flat': True,  # No intenta procesar cada video del playlist
        'force_generic_extractor': True,  # Evita usar extractores específicos que fallan
    }

    with yt.YoutubeDL(ydl_opts) as ydl:
        try:
            info = ydl.extract_info(url, download=False)
            album_name = info.get('album') or info.get('playlist_title') or info.get('title')
            if album_name:
                album_name = album_name.replace("Album - ", "").strip()
            print(album_name)
            return album_name
        except Exception as e:
            print(f"⚠️ Error obteniendo nombre del álbum: {e}")
            return None
        
obtener_nombre_album(url)