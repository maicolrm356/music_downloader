import openpyxl
from openpyxl import load_workbook, Workbook
from openpyxl.utils import get_column_letter
def cargarExcelYaccederAHojaDeseada(file, sheet_obj):
    book = load_workbook(file)
    sheet = book[sheet_obj]
    columnas = sheet.max_column
    filas = sheet.max_row
    sheet.iter_rows
    sheet.iter_cols
    # recorreExcel(sheet)
    print(filas)
    print(columnas)
    # obtenerNombresDeAlbumes(sheet, columnas)
    contarFilasEnColumnas(sheet, columnas, filas)

def recorreExcel(sheet):
    max_row = sheet.max_row
    max_column = sheet.max_column
    
    for column in range(1, max_column + 1):
        print(f"Columna: {column}")


    for row in range(0, sheet.max_row):
        for col in sheet.iter_cols(1, sheet.max_column):
            print(col[row].value)

# book = xls.add
def obtenerNombresDeAlbumes(sheet, columnas):
    albums = []
    for row in sheet.iter_rows(min_row=1, max_row=1, min_col=1, max_col=columnas):
        for celda in row:
            albums.append(celda.value)
    print(albums)
    return albums

# def obtenerNombreAlbumYSusCanciones(sheet, columnas):
    # for row in sheet.iter_cols(min_col=1, min_col=columnas, min_row=1,)

def contarFilasEnColumnas(sheet,columnas, filas):
    for col in range(1, columnas +1):
        contador = 0
        for fila in range(1, filas + 1):
            celda = sheet.cell(row=fila, column=col)
            valor = celda.value
            if fila == 1:
                print(valor)
            if valor is not None:
                contador += 1
        print(f"Columna {col} tiene {contador} filas con datos.")
    print(contador)

# def crearCarpetas(nombre_album):
#     path = "C:\musica"


# ingresar a la hoja de excel deseada
# sheet = xls['PXNDX']
# # sheets = xls.read_worksheets()
# print(f"numero de filas: {sheet.max_row} \nnumero de columnas: {sheet.max_column}")

# sheet.columns
# cargarExcelYaccederAHojaDeseada("./music.xlsx", "PXNDX")


# Carga el libro de trabajo
try:
    workbook = load_workbook('./music.xlsx')
except FileNotFoundError:
    print("El archivo no se encontró.")
    exit()

# Obtiene los nombres de las hojas en una lista
sheet_names = workbook.sheetnames

# Imprime los nombres de las hojas
print(sheet_names)

# Ejemplo de cómo iterar sobre las hojas
for sheet_name in sheet_names:
    print(f"Nombre de la hoja: {sheet_name}")
    # Puedes realizar otras acciones con cada hoja aquí
    # sheet = workbook[sheet_name]